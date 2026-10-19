"""Stock-movement workbook: parsing, discrepancy detection, unit confirmation, safe import.

Uses a synthetic workbook with the same layout as CFC's "oct 2026 stock movement.xlsx"
(real business data is never committed to the repository)."""
import io

from openpyxl import Workbook

from app import workbook
from conftest import login, open_day


def make_workbook(break_carry=False, dup=False):
    wb = Workbook()
    wb.remove(wb.active)
    data = {  # sheet: [(item, opening, additions, closing, system, note)]
        "1st": [("raw goat meat", 19.62, None, "=13.82", 4.2, None), ("chapati", 6, 20, 12, 7, "overstayed"),
                ("sodas", 345, None, 335, 10, None), ("samosa", "=15+20+25", "=32", "=37+16+19", 20, None)],
        "2nd": [("raw goat meat", "=13.82", None, 5.06, 8.75, None), ("chapati", 12, 22, 0, 20, None),
                ("sodas", 335 if not break_carry else 330, None, 314, 22, None), ("samosa", "=37+16+19", None, 60, 12, None)],
        "4th-5th": [("raw goat meat", 5.06, 18.6, 10.32, 12.63, None), ("chapati", 0, 43, 12, 32, None),
                    ("sodas", 314, None, 267, 47, None), ("samosa", 60, 18, 48, 30, None)],
    }
    for name, rows in data.items():
        ws = wb.create_sheet(name)
        for col, h in zip("CDEFGH", ["opening", "additions", "closing", "actual sales", "system sales", "variance"]):
            ws[f"{col}2"] = h
        for i, (item, o, a, c, sysv, note) in enumerate(rows + ([rows[0]] if dup and name == "1st" else []), start=3):
            ws[f"B{i}"], ws[f"C{i}"], ws[f"D{i}"], ws[f"E{i}"] = item, o, a, c
            ws[f"F{i}"] = f"=C{i}+D{i}-E{i}"
            ws[f"G{i}"] = sysv
            ws[f"H{i}"] = f"=G{i}-F{i}"
            if note:
                ws[f"I{i}"] = note
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def test_parse_reproduces_workbook_formulas():
    r = workbook.parse(make_workbook(), "oct 2026 stock movement.xlsx")
    assert (r["year"], r["month"]) == (2026, 10)
    first = {x["key"]: x for x in r["sheets"][0]["rows"]}
    assert abs(first["raw goat meat"]["actual"] - 5.8) < 1e-9 and abs(first["raw goat meat"]["variance"] + 1.6) < 1e-9
    assert first["samosa"]["opening"] == 60 and first["samosa"]["closing"] == 72
    codes = {i["code"] for i in r["issues"]}
    assert {"multi_day", "missing_day", "split_count", "note"} <= codes
    assert "carry_over_break" not in codes
    units = {i["key"]: i["suggested_unit"] for i in r["items"]}
    assert units == {"raw goat meat": "kg", "chapati": "pcs", "sodas": "pcs", "samosa": "pcs"}


def test_detects_carry_over_break_and_duplicates():
    r = workbook.parse(make_workbook(break_carry=True, dup=True), "oct 2026.xlsx")
    codes = [i["code"] for i in r["issues"]]
    assert "carry_over_break" in codes and "duplicate_item" in codes


def test_import_needs_confirmed_units_and_is_safe(cl):
    M = login(cl, "2222")
    data = make_workbook()
    f = {"file": ("oct 2026 stock movement.xlsx", data)}
    pv = cl.post("/api/import/stock-workbook", headers=M, files=f).json()
    assert pv["commit"] is False and len(pv["items"]) == 4
    # recipe importer refuses this file type instead of misreading it
    assert "stock-movement" in cl.post("/api/excel", headers=M, files=f).json()["detail"]
    r = cl.post("/api/import/stock-workbook?commit=true", headers=M, files=f, data={"mapping": "{}"})
    assert r.status_code == 400 and "choose the unit" in r.json()["detail"]
    mapping = '{"raw goat meat": {"unit": "kg"}, "chapati": {"unit": "pcs"}, "sodas": {"unit": "pcs"}, "samosa": {"unit": "pcs"}}'
    r = cl.post("/api/import/stock-workbook?commit=true", headers=M, files=f, data={"mapping": mapping})
    assert r.status_code == 200, r.text
    assert r.json()["days"] == 3 and len(r.json()["new_ingredients"]) == 4
    # same file twice -> refused, nothing duplicated
    r2 = cl.post("/api/import/stock-workbook?commit=true", headers=M, files=f, data={"mapping": mapping})
    assert r2.status_code == 400 and "already imported" in r2.json()["detail"]
    days = cl.get("/api/days", headers=M).json()
    assert [d["business_date"] for d in days] == ["2026-10-04", "2026-10-02", "2026-10-01"] and days[0]["covers_to"] == "2026-10-05"
    rep = cl.get(f"/api/reports/day/{days[-1]['id']}", headers=M).json()
    goat = next(l for l in rep["stock"] if l["name"] == "raw goat meat")
    assert goat["unit"] == "g" and goat["opening"] == 19620 and abs(goat["variance"] + 1600) < 1e-6  # same as sheet: -1.6 kg
    chap = next(l for l in rep["stock"] if l["name"] == "chapati")
    assert chap["variance"] == -7 and "overstayed" in chap["reason"]
    # last workbook closing becomes the next opening (carry-over)
    prev = {r["name"]: r["carry_over"] for r in cl.get("/api/day/opening-preview", headers=M).json()["rows"]}
    assert prev["raw goat meat"] == 10320 and prev["sodas"] == 267
    open_day_body = {"business_date": "2026-10-03"}
    assert cl.post("/api/day/open", headers=M, json=open_day_body).status_code == 400  # inside imported history


def test_import_refused_when_app_already_has_later_days(cl):
    M = login(cl, "2222")
    open_day(cl, M, date="2026-10-02")
    f = {"file": ("oct 2026 stock movement.xlsx", make_workbook())}
    mapping = '{"raw goat meat": {"unit": "kg"}, "chapati": {"unit": "pcs"}, "sodas": {"unit": "pcs"}, "samosa": {"unit": "pcs"}}'
    r = cl.post("/api/import/stock-workbook?commit=true", headers=M, files=f, data={"mapping": mapping})
    assert r.status_code == 400 and "open day" in r.json()["detail"].lower()


def test_unit_must_match_existing_item(cl):
    M = login(cl, "2222")
    f = {"file": ("oct 2026 stock movement.xlsx", make_workbook())}
    cl.post("/api/ingredients", headers=M, json={"name": "sodas", "unit": "ml"})
    mapping = '{"raw goat meat": {"unit": "kg"}, "chapati": {"unit": "pcs"}, "sodas": {"unit": "pcs"}, "samosa": {"unit": "pcs"}}'
    r = cl.post("/api/import/stock-workbook?commit=true", headers=M, files=f, data={"mapping": mapping})
    assert r.status_code == 400 and "measured in ml" in r.json()["detail"]
    assert cl.get("/api/days", headers=M).json() == []  # nothing saved
