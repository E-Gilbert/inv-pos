"""Reader for CFC's daily stock-movement workbook (e.g. "oct 2026 stock movement.xlsx").

Layout of every sheet (one sheet per business day, named "1st", "2nd", "4th-5th", ...):

    B: item | C: opening | D: additions | E: closing | F: actual sales | G: system sales | H: variance | I: note

Workbook formulas (reproduced exactly here):

    actual sales (F) = opening + additions - closing          physical usage
    variance     (H) = system sales - actual sales            negative = more left the store than the POS sold

That is the same quantity INV-POS calls the closing variance:

    expected = opening + delivered - sold - waste;   variance = counted - expected
             = (sold + waste) - (opening + delivered - counted)  ==  system - actual   (waste = 0 here)

so imported history lines up with app-generated days without sign changes.

This module only READS the workbook (it never writes to it) and returns plain data plus a
list of discrepancies. `import_history` (in importer.py) commits it to the database.
"""
import hashlib
import io
import re
from collections import OrderedDict
from datetime import date

from openpyxl import load_workbook

MONTHS = {m: i for i, m in enumerate(
    ["jan", "feb", "mar", "apr", "may", "jun", "jul", "aug", "sep", "oct", "nov", "dec"], start=1)}
DAY_RE = re.compile(r"(\d{1,2})\s*(?:st|nd|rd|th)?", re.I)
COLS = {"opening": "opening", "additions": "additions", "closing": "closing",
        "actual sales": "actual", "system sales": "system", "variance": "variance"}
EPS = 1e-6

# Units the workbook can be in, and how they map to the app's base units.
WORKBOOK_UNITS = {"kg": ("g", 1000), "g": ("g", 1), "l": ("ml", 1000), "ml": ("ml", 1), "pcs": ("pcs", 1)}


def norm(name):
    return re.sub(r"\s+", " ", str(name or "").strip()).lower()


def period_from_filename(filename):
    """'oct 2026 stock movement.xlsx' -> (2026, 10). None if it can't tell."""
    f = (filename or "").lower()
    m = re.search(r"\b(jan|feb|mar|apr|may|jun|jul|aug|sep|oct|nov|dec)[a-z]*[\s_-]*(20\d\d)\b", f)
    if m:
        return int(m.group(2)), MONTHS[m.group(1)]
    m = re.search(r"\b(20\d\d)[\s_-]*(0?[1-9]|1[0-2])\b", f)
    return (int(m.group(1)), int(m.group(2))) if m else None


def _days_from_sheet(name):
    """'1st' -> [1]; '4th-5th' -> [4, 5]. Returns [] if the name has no day."""
    nums = [int(n) for n in DAY_RE.findall(name)]
    nums = [n for n in nums if 1 <= n <= 31]
    if len(nums) >= 2:
        return list(range(nums[0], nums[-1] + 1)) if nums[-1] >= nums[0] else []
    return nums


def _num(v):
    if v is None or v == "":
        return None
    if isinstance(v, (int, float)):
        return float(v)
    try:
        return float(str(v).replace(",", "").strip())
    except ValueError:
        return "bad"


def _parts(formula):
    """'=15+20+25' -> 3 (a count split over several places). Plain '=13.82' -> 1."""
    if isinstance(formula, str) and formula.startswith("="):
        body = formula[1:]
        if re.fullmatch(r"[\d.\s+]+", body):
            return len([p for p in body.split("+") if p.strip()])
    return 0


def _literal_sum(formula):
    body = formula[1:].replace(" ", "")
    if re.fullmatch(r"\d+(\.\d+)?(\+\d+(\.\d+)?)*", body):
        return float(sum(float(p) for p in body.split("+")))
    return None


def suggest_unit(values, name):
    """A SUGGESTION only (the user must confirm). Whole numbers everywhere -> pieces;
    otherwise a unit only when the name makes it obvious."""
    n = norm(name)
    nums = [v for v in values if isinstance(v, float)]
    if nums and all(abs(v - round(v)) < EPS for v in nums):
        return "pcs"
    if any(k in n for k in ("oil", "milk", "juice")):
        return "L"
    if any(k in n for k in ("meat", "goat", "beef", "rice", "flour", "sugar", "steak")):
        return "kg"
    return None


def parse(data: bytes, filename="", year=None, month=None):
    wb_f = load_workbook(io.BytesIO(data), data_only=False)
    wb_v = load_workbook(io.BytesIO(data), data_only=True)
    if not (year and month):
        p = period_from_filename(filename)
        if p:
            year, month = p
    issues = []

    def issue(level, sheet, item, msg, code):
        issues.append({"level": level, "sheet": sheet, "item": item, "message": msg, "code": code})

    sheets = []
    for ws in wb_f.worksheets:
        wv = wb_v[ws.title]
        header_row, cols = None, {}
        for row in ws.iter_rows(min_row=1, max_row=min(ws.max_row, 10)):
            found = {COLS[norm(c.value)]: c.column for c in row if isinstance(c.value, str) and norm(c.value) in COLS}
            if "opening" in found and "closing" in found:
                header_row, cols = row[0].row, found
                break
        if not header_row:
            issue("warning", ws.title, None, "No 'opening / closing' header found; sheet skipped.", "no_header")
            continue
        item_col = min(cols.values()) - 1
        note_col = max(cols.values()) + 1
        days = _days_from_sheet(ws.title)
        if not days:
            issue("error", ws.title, None, "Sheet name has no day number; can't date it.", "no_date")
        dates = []
        if days and year and month:
            try:
                dates = [date(year, month, d).isoformat() for d in days]
            except ValueError:
                issue("error", ws.title, None, "Sheet day is not a valid date for the month.", "bad_date")
        if len(days) > 1:
            issue("warning", ws.title, None,
                  f"One sheet covers {len(days)} days ({ws.title}); it is imported as a single period "
                  f"dated {dates[0] if dates else '?'} so daily figures for those days can't be separated.",
                  "multi_day")
        rows, seen = [], set()
        for r in range(header_row + 1, ws.max_row + 1):
            name = ws.cell(r, item_col).value
            if name in (None, ""):
                continue
            key = norm(name)
            if key in seen:
                issue("error", ws.title, str(name), "Item appears twice on the same sheet.", "duplicate_item")
                continue
            seen.add(key)
            vals, raw, parts = {}, {}, {}
            for k, col in cols.items():
                f = ws.cell(r, col).value
                v = wv.cell(r, col).value
                raw[k] = f
                parts[k] = _parts(f)
                if isinstance(f, str) and f.startswith("=") and v is None:
                    # No value saved by Excel (file written by another tool). Simple typed sums
                    # like =15+20+25 are evaluated here; the actual/variance columns are recomputed anyway.
                    lit = _literal_sum(f)
                    if lit is not None:
                        v = lit
                    elif k not in ("actual", "variance"):
                        issue("error", ws.title, str(name), f"{k} is a formula with no saved value; open and save it in Excel.", "no_cached_value")
                elif not isinstance(f, str) or not f.startswith("="):
                    v = f if v is None else v
                vals[k] = _num(v)
            if any(vals[k] == "bad" for k in vals):
                issue("error", ws.title, str(name), "A cell that should be a number contains text.", "text_in_number")
                continue
            note = wv.cell(r, note_col).value
            o, a, c, s = (vals.get("opening") or 0.0), (vals.get("additions") or 0.0), vals.get("closing"), vals.get("system")
            if vals.get("opening") is None:
                issue("warning", ws.title, str(name), "Opening is blank (treated as 0).", "blank_opening")
            if c is None:
                issue("error", ws.title, str(name), "Closing count is blank.", "blank_closing")
                continue
            actual = round(o + a - c, 6)
            variance = round((s or 0) - actual, 6)
            if vals.get("actual") is not None and abs(vals["actual"] - actual) > 1e-4:
                issue("error", ws.title, str(name),
                      f"'actual sales' cell shows {vals['actual']:g} but opening+additions-closing = {actual:g}.", "formula_mismatch")
            if vals.get("variance") is not None and abs(vals["variance"] - variance) > 1e-4:
                issue("error", ws.title, str(name),
                      f"'variance' cell shows {vals['variance']:g} but system-actual = {variance:g}.", "formula_mismatch")
            if actual < -EPS:
                issue("error", ws.title, str(name),
                      f"Closing ({c:g}) is more than opening + additions ({o + a:g}): a count or a delivery is wrong.", "negative_usage")
            if s is None:
                issue("warning", ws.title, str(name), "System sales blank.", "blank_system")
            if variance > EPS:
                issue("warning", ws.title, str(name),
                      f"POS sold {s:g} but only {actual:g} left the store (+{variance:g}). Physically impossible unless "
                      "a delivery was not written down, the count is wrong, or another item was rung up instead.",
                      "positive_variance")
            for k in ("opening", "additions", "closing"):
                if parts[k] > 1:
                    issue("info", ws.title, str(name), f"{k} is the sum of {parts[k]} separate counts ({raw[k]}).", "split_count")
            rows.append({
                "item": str(name).strip(), "key": key, "opening": o, "additions": a, "closing": c,
                "actual": actual, "system": s, "variance": variance, "note": str(note).strip() if note else None,
            })
            if note:
                issue("info", ws.title, str(name), f"Note on sheet: {note}", "note")
        sheets.append({"name": ws.title, "days": days, "dates": dates, "rows": rows})

    # Order chronologically, then check gaps and carry-over between consecutive sheets.
    sheets.sort(key=lambda s: (s["days"][0] if s["days"] else 99))
    covered = [d for s in sheets for d in s["days"]]
    for s in sheets:
        if s["days"] and covered.count(s["days"][0]) > 1:
            issue("error", s["name"], None, "Another sheet covers the same day.", "overlap")
    if covered:
        missing = [d for d in range(min(covered), max(covered) + 1) if d not in covered]
        for d in missing:
            issue("warning", None, None, f"No sheet for day {d} of the month: stock movement for that day is unrecorded.", "missing_day")
    for prev, cur in zip(sheets, sheets[1:]):
        pc = {r["key"]: r for r in prev["rows"]}
        for r in cur["rows"]:
            p = pc.get(r["key"])
            if p is None:
                issue("warning", cur["name"], r["item"], f"Not on the previous sheet ({prev['name']}).", "new_item")
            elif abs(p["closing"] - r["opening"]) > 1e-6:
                issue("error", cur["name"], r["item"],
                      f"Opening {r['opening']:g} doesn't match previous closing {p['closing']:g} on '{prev['name']}' "
                      f"(difference {r['opening'] - p['closing']:+g}).", "carry_over_break")
        for k in pc:
            if k not in {r["key"] for r in cur["rows"]}:
                issue("warning", cur["name"], pc[k]["item"], f"On '{prev['name']}' but missing here.", "dropped_item")

    # Per-item summary across the period.
    items = OrderedDict()
    for s in sheets:
        for r in s["rows"]:
            it = items.setdefault(r["key"], {"item": r["item"], "values": [], "actual": 0.0, "system": 0.0,
                                             "variance": 0.0, "days_off": 0, "additions": 0.0})
            it["values"] += [r["opening"], r["additions"], r["closing"]] + ([r["system"]] if r["system"] is not None else [])
            it["actual"] += r["actual"]
            it["system"] += r["system"] or 0
            it["variance"] += r["variance"]
            it["additions"] += r["additions"]
            it["days_off"] += abs(r["variance"]) > 1e-6
            it["last_closing"] = r["closing"]
    summary = []
    for k, it in items.items():
        summary.append({
            "key": k, "item": it["item"], "suggested_unit": suggest_unit(it["values"], it["item"]),
            "total_actual": round(it["actual"], 4), "total_system": round(it["system"], 4),
            "total_variance": round(it["variance"], 4), "days_with_variance": it["days_off"],
            "total_additions": round(it["additions"], 4), "last_closing": it["last_closing"],
        })

    return {
        "year": year, "month": month, "sha256": hashlib.sha256(data).hexdigest(),
        "sheets": sheets, "items": summary, "issues": issues,
        "counts": {lvl: sum(1 for i in issues if i["level"] == lvl) for lvl in ("error", "warning", "info")},
    }


def is_stock_movement_workbook(data: bytes):
    """True if the first sheet has the opening/closing header (vs. the recipe template)."""
    try:
        ws = load_workbook(io.BytesIO(data), read_only=True).worksheets[0]
        for row in ws.iter_rows(min_row=1, max_row=10, values_only=True):
            names = {norm(v) for v in row if isinstance(v, str)}
            if {"opening", "closing"} <= names:
                return True
    except Exception:
        return False
    return False
