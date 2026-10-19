"""Excel import/export for ingredients, menu and recipes.

Template (two sheets):
  Ingredients: name | unit (g/ml/pcs) | cost_per_unit (KES) | par_level | count_daily (yes/no)
  Recipes:     menu_item | category | price | ingredient | qty_per_portion

One row in Recipes = one ingredient in one dish. A dish with 4 ingredients has 4 rows;
category and price only need to be filled on the dish's first row.

CFC's own sheet will probably look different; once we have it, either fill this
template from it or add a small adapter function here.
"""
import csv
import io
import json

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

class ImportErrors(ValueError):
    """All row errors at once, so the user can fix the whole file in one go."""
    def __init__(self, errors):
        self.errors = errors
        super().__init__("\n".join(errors[:40]) + (f"\n…and {len(errors) - 40} more" if len(errors) > 40 else ""))


ING_HEADERS = ["name", "unit", "cost_per_unit", "par_level", "count_daily"]
REC_HEADERS = ["menu_item", "category", "price", "ingredient", "qty_per_portion"]
UNIT_ALIASES = {
    "g": "g", "gram": "g", "grams": "g", "gm": "g", "gms": "g",
    "ml": "ml", "millilitre": "ml", "milliliter": "ml",
    "pcs": "pcs", "pc": "pcs", "piece": "pcs", "pieces": "pcs", "each": "pcs", "unit": "pcs",
}
# Convert bigger units to the base unit so staff can type what they're used to.
SCALE_UNITS = {"kg": ("g", 1000), "kgs": ("g", 1000), "l": ("ml", 1000), "litre": ("ml", 1000),
               "liter": ("ml", 1000), "ltr": ("ml", 1000)}


def _rows(ws):
    it = ws.iter_rows(values_only=True)
    try:
        header = [str(h).strip().lower() if h is not None else "" for h in next(it)]
    except StopIteration:
        return []
    out = []
    for n, row in enumerate(it, start=2):
        if row is None or all(v in (None, "") for v in row):
            continue
        out.append((n, {header[i]: row[i] for i in range(min(len(header), len(row))) if header[i]}))
    return out


def _yes(v):
    return str(v).strip().lower() in ("1", "yes", "y", "true", "daily") if v not in (None, "") else True


def import_workbook(c, data: bytes):
    wb = load_workbook(io.BytesIO(data), data_only=True)
    sheets = {s.lower(): wb[s] for s in wb.sheetnames}
    if "ingredients" not in sheets or "recipes" not in sheets:
        raise ValueError("The workbook needs two sheets named 'Ingredients' and 'Recipes'. Download the template.")

    errors, stats = [], {"ingredients": 0, "new_ingredients": 0, "updated": 0, "menu_items": 0, "recipe_lines": 0}
    seen_ings, seen_pairs, changes = set(), set(), []

    for n, r in _rows(sheets["ingredients"]):
        name = str(r.get("name") or "").strip()
        unit_raw = str(r.get("unit") or "").strip().lower()
        if not name:
            errors.append(f"Ingredients row {n}: name is empty.")
            continue
        scale = 1
        if unit_raw in SCALE_UNITS:
            unit, scale = SCALE_UNITS[unit_raw]
        else:
            unit = UNIT_ALIASES.get(unit_raw)
        if not unit:
            errors.append(f"Ingredients row {n} ({name}): unit '{unit_raw}' must be g, ml or pcs (kg/L are converted).")
            continue
        try:
            # cost given per kg/L gets converted to cost per g/ml
            cost = float(r.get("cost_per_unit") or 0) / scale
            par = float(r.get("par_level") or 0) * scale
        except (TypeError, ValueError):
            errors.append(f"Ingredients row {n} ({name}): cost and par level must be numbers.")
            continue
        if name.lower() in seen_ings:
            errors.append(f"Ingredients row {n} ({name}): listed twice in the sheet.")
            continue
        seen_ings.add(name.lower())
        if cost < 0 or par < 0:
            errors.append(f"Ingredients row {n} ({name}): cost and par level can't be negative.")
            continue
        old = c.execute("SELECT * FROM ingredients WHERE name = ?", (name,)).fetchone()
        if old and old["unit"] != unit and c.execute(
                "SELECT 1 FROM stock_movements WHERE ingredient_id = ? UNION SELECT 1 FROM day_stock WHERE ingredient_id = ? LIMIT 1",
                (old["id"], old["id"])).fetchone():
            errors.append(f"Ingredients row {n} ({name}): unit can't change from {old['unit']} to {unit}; it already has stock history.")
            continue
        if old:
            stats["updated"] += 1
            if abs((old["cost_per_unit"] or 0) - cost) > 1e-9:
                changes.append({"ingredient": name, "field": "cost_per_unit", "old": old["cost_per_unit"], "new": cost})
        else:
            stats["new_ingredients"] += 1
        c.execute(
            """INSERT INTO ingredients(name, unit, cost_per_unit, par_level, count_daily) VALUES (?, ?, ?, ?, ?)
               ON CONFLICT(name) DO UPDATE SET unit = excluded.unit, cost_per_unit = excluded.cost_per_unit,
                   par_level = excluded.par_level, count_daily = excluded.count_daily, active = 1""",
            (name, unit, cost, par, 1 if _yes(r.get("count_daily")) else 0),
        )
        stats["ingredients"] += 1

    current_item, seen_items = None, set()
    for n, r in _rows(sheets["recipes"]):
        item_name = str(r.get("menu_item") or "").strip() or current_item
        if not item_name:
            errors.append(f"Recipes row {n}: menu_item is empty.")
            continue
        current_item = item_name
        mi = c.execute("SELECT id FROM menu_items WHERE name = ?", (item_name,)).fetchone()

        cat, price = r.get("category"), r.get("price")
        if cat not in (None, "") or price not in (None, ""):
            cat_id = None
            if cat not in (None, ""):
                c.execute("INSERT OR IGNORE INTO categories(name) VALUES (?)", (str(cat).strip(),))
                cat_id = c.execute("SELECT id FROM categories WHERE name = ?", (str(cat).strip(),)).fetchone()["id"]
            try:
                price_f = float(price) if price not in (None, "") else None
            except (TypeError, ValueError):
                errors.append(f"Recipes row {n} ({item_name}): price must be a number.")
                continue
            if price_f is not None and price_f < 0:
                errors.append(f"Recipes row {n} ({item_name}): price can't be negative.")
                continue
            if mi:
                if cat_id:
                    c.execute("UPDATE menu_items SET category_id = ? WHERE id = ?", (cat_id, mi["id"]))
                old_price = c.execute("SELECT price FROM menu_items WHERE id = ?", (mi["id"],)).fetchone()["price"]
                if price_f is not None and abs(old_price - price_f) > 1e-9:
                    changes.append({"menu_item": item_name, "field": "price", "old": old_price, "new": price_f})
                if price_f is not None:
                    c.execute("UPDATE menu_items SET price = ?, active = 1 WHERE id = ?", (price_f, mi["id"]))
            else:
                if price_f is None:
                    errors.append(f"Recipes row {n} ({item_name}): new menu item needs a price.")
                    continue
                c.execute("INSERT INTO menu_items(name, category_id, price) VALUES (?, ?, ?)", (item_name, cat_id, price_f))
            mi = c.execute("SELECT id FROM menu_items WHERE name = ?", (item_name,)).fetchone()
        if not mi:
            errors.append(f"Recipes row {n}: '{item_name}' isn't on the menu yet; give it a price.")
            continue

        if mi["id"] not in seen_items:
            # The sheet is the source of truth: replace the dish's whole recipe.
            c.execute("DELETE FROM recipe_lines WHERE menu_item_id = ?", (mi["id"],))
            seen_items.add(mi["id"])
            stats["menu_items"] += 1

        ing_name = str(r.get("ingredient") or "").strip()
        if not ing_name:
            continue  # a dish with no tracked ingredients (allowed)
        ing = c.execute("SELECT id, unit FROM ingredients WHERE name = ?", (ing_name,)).fetchone()
        if not ing:
            errors.append(f"Recipes row {n} ({item_name}): ingredient '{ing_name}' isn't in the Ingredients sheet.")
            continue
        if (mi["id"], ing["id"]) in seen_pairs:
            errors.append(f"Recipes row {n} ({item_name} / {ing_name}): this ingredient is already listed for this dish; combine the rows.")
            continue
        seen_pairs.add((mi["id"], ing["id"]))
        try:
            qty = float(r.get("qty_per_portion"))
            if not qty > 0 or qty != qty:
                raise ValueError
        except (TypeError, ValueError):
            errors.append(f"Recipes row {n} ({item_name} / {ing_name}): qty_per_portion must be more than 0.")
            continue
        c.execute(
            """INSERT INTO recipe_lines(menu_item_id, ingredient_id, qty) VALUES (?, ?, ?)
               ON CONFLICT(menu_item_id, ingredient_id) DO UPDATE SET qty = excluded.qty""",
            (mi["id"], ing["id"], qty),
        )
        stats["recipe_lines"] += 1

    if errors:
        # Nothing is saved unless the whole sheet is clean.
        raise ImportErrors(errors)
    stats["changes"] = changes
    return stats


def _style_header(ws, headers, widths):
    ws.append(headers)
    for i, cell in enumerate(ws[1]):
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="7A2E1D")
        ws.column_dimensions[cell.column_letter].width = widths[i]
    ws.freeze_panes = "A2"


def export_workbook(c) -> bytes:
    """Current data in template format. Empty database -> blank template with examples."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Ingredients"
    _style_header(ws, ING_HEADERS, [28, 8, 16, 12, 13])
    dv = DataValidation(type="list", formula1='"g,ml,pcs,kg,L"', allow_blank=False)
    ws.add_data_validation(dv)
    dv.add("B2:B1000")
    ings = c.execute("SELECT * FROM ingredients WHERE active = 1 ORDER BY name").fetchall()
    for i in ings:
        ws.append([i["name"], i["unit"], i["cost_per_unit"], i["par_level"], "yes" if i["count_daily"] else "no"])
    if not ings:
        ws.append(["Basmati rice", "kg", 250, 10, "yes"])
        ws.append(["Chicken (raw)", "kg", 550, 8, "yes"])

    wr = wb.create_sheet("Recipes")
    _style_header(wr, REC_HEADERS, [26, 16, 10, 26, 16])
    rows = c.execute(
        """SELECT m.name, c.name AS cat, m.price, i.name AS ing, rl.qty, m.id
           FROM menu_items m LEFT JOIN categories c ON c.id = m.category_id
           LEFT JOIN recipe_lines rl ON rl.menu_item_id = m.id
           LEFT JOIN ingredients i ON i.id = rl.ingredient_id
           WHERE m.active = 1 ORDER BY c.sort, c.name, m.name, i.name"""
    ).fetchall()
    last = None
    for r in rows:
        first = r["id"] != last
        wr.append([r["name"], r["cat"] if first else None, r["price"] if first else None, r["ing"], r["qty"]])
        last = r["id"]
    if not rows:
        wr.append(["Chicken Mandi", "Mandi", 950, "Basmati rice", 250])
        wr.append(["Chicken Mandi", None, None, "Chicken (raw)", 400])

    notes = wb.create_sheet("How to fill")
    for line in [
        "INGREDIENTS sheet: one row per raw material.",
        "  unit: g, ml or pcs. You may type kg or L and the system converts (cost per kg -> per g).",
        "  count_daily: yes = must be counted every closing (meat, rice, oil). no = weekly items.",
        "  par_level: when stock falls below this, it shows as 'reorder'. Same unit as you typed.",
        "",
        "RECIPES sheet: one row per ingredient per dish.",
        "  qty_per_portion is in the ingredient's BASE unit (g, ml, pcs), for ONE plate.",
        "  Category and price only need to be on the first row of each dish.",
        "  Use the RAW weight (e.g. raw chicken before cooking), so cooking loss is included.",
        "",
        "Importing replaces each listed dish's recipe. If any row has an error, nothing is saved.",
    ]:
        notes.append([line])
    notes.column_dimensions["A"].width = 100

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


# ---------------------------------------------------------------- CSV (ingredients)

def import_ingredients_csv(c, data: bytes):
    """Same columns as the Ingredients sheet. Wrapped into the workbook importer's rules."""
    try:
        text = data.decode("utf-8-sig")
    except UnicodeDecodeError:
        text = data.decode("latin-1")
    rows = list(csv.reader(io.StringIO(text)))
    if not rows:
        raise ValueError("The CSV file is empty.")
    wb = Workbook()
    ws = wb.active
    ws.title = "Ingredients"
    for r in rows:
        ws.append(r)
    wb.create_sheet("Recipes").append(REC_HEADERS)
    buf = io.BytesIO()
    wb.save(buf)
    return import_workbook(c, buf.getvalue())


def to_csv(rows, columns):
    """rows: list of dicts. Cells starting with = + - @ are prefixed so Excel never runs them as formulas."""
    buf = io.StringIO()
    w = csv.writer(buf)
    w.writerow([h for _, h in columns])
    for r in rows:
        out = []
        for key, _ in columns:
            v = r.get(key)
            if isinstance(v, str) and v[:1] in ("=", "+", "-", "@", "\t", "\r"):
                v = "'" + v
            out.append("" if v is None else v)
        w.writerow(out)
    return ("﻿" + buf.getvalue()).encode("utf-8")  # BOM so Excel opens UTF-8 correctly


# ---------------------------------------------------------------- stock-movement workbook (history)

def import_history(c, parsed, mapping, staff_id, filename=None):
    """Commit a parsed stock-movement workbook (see workbook.py) as closed, read-only history days.

    mapping: {item_key: {"unit": "kg|g|L|ml|pcs", "ingredient_id": optional existing id}}
    Nothing is guessed: every item needs a confirmed unit. All-or-nothing.
    """
    from .db import audit, now
    from .workbook import WORKBOOK_UNITS

    errors = [f"{i['sheet'] or ''} {i['item'] or ''}: {i['message']}".strip()
              for i in parsed["issues"] if i["level"] == "error"]
    if errors:
        raise ImportErrors(["The workbook has errors that must be fixed first:"] + errors)
    if not (parsed.get("year") and parsed.get("month")):
        raise ValueError("Which month and year is this workbook? (Name it like 'oct 2026 …' or choose the month.)")
    if c.execute("SELECT 1 FROM import_batches WHERE kind = 'stock_history' AND sha256 = ?", (parsed["sha256"],)).fetchone():
        raise ValueError("This exact workbook was already imported. Nothing was changed.")
    if c.execute("SELECT 1 FROM days WHERE status = 'open'").fetchone():
        raise ValueError("Close the open day before importing history.")
    sheets = [s for s in parsed["sheets"] if s["rows"] and s["dates"]]
    if not sheets:
        raise ValueError("No dated sheets with data found.")
    first = sheets[0]["dates"][0]
    clash = c.execute("SELECT business_date FROM days WHERE COALESCE(covers_to, business_date) >= ? ORDER BY business_date LIMIT 1",
                      (first,)).fetchone()
    if clash:
        raise ValueError(f"The system already has a day on or after {first} ({clash['business_date']}). History can only be "
                         "imported for dates before everything already recorded, otherwise opening stock would be wrong.")

    resolved, errs = {}, []
    mapping = {k.lower(): v for k, v in (mapping or {}).items()}
    for it in parsed["items"]:
        m = mapping.get(it["key"]) or {}
        unit = str(m.get("unit") or "").strip().lower()
        if unit not in WORKBOOK_UNITS:
            errs.append(f"{it['item']}: choose the unit the workbook counts it in (kg, g, L, ml or pcs).")
            continue
        base, scale = WORKBOOK_UNITS[unit]
        ing = None
        if m.get("ingredient_id"):
            ing = c.execute("SELECT * FROM ingredients WHERE id = ?", (int(m["ingredient_id"]),)).fetchone()
            if not ing:
                errs.append(f"{it['item']}: the chosen existing item no longer exists.")
                continue
        else:
            ing = c.execute("SELECT * FROM ingredients WHERE name = ?", (it["item"],)).fetchone()
        if ing and ing["unit"] != base:
            errs.append(f"{it['item']}: '{ing['name']}' is measured in {ing['unit']} but you chose {unit}.")
            continue
        if base == "pcs":
            fractional = [s_["name"] for s_ in parsed["sheets"] for row in s_["rows"] if row["key"] == it["key"]
                          and any(v is not None and abs(v - round(v)) > 1e-9 for v in (row["opening"], row["additions"], row["closing"], row["system"]))]
            if fractional:
                errs.append(f"{it['item']}: has fractional quantities (e.g. sheet '{fractional[0]}'), so it can't be counted in pieces. "
                            "Choose kg/g/L/ml, or correct the sheet.")
                continue
        resolved[it["key"]] = {"ing": ing, "base": base, "scale": scale, "name": it["item"], "unit": unit}
    if errs:
        raise ImportErrors(errs)

    batch = c.execute("INSERT INTO import_batches(kind, filename, sha256, staff_id, created_at) VALUES ('stock_history', ?, ?, ?, ?)",
                      (filename, parsed["sha256"], staff_id, now())).lastrowid
    created = []
    for r in resolved.values():
        if r["ing"] is None:
            iid = c.execute("INSERT INTO ingredients(name, unit, cost_per_unit, par_level, count_daily) VALUES (?, ?, 0, 0, 1)",
                            (r["name"], r["base"])).lastrowid
            r["ing"] = c.execute("SELECT * FROM ingredients WHERE id = ?", (iid,)).fetchone()
            created.append(r["name"])
    stats = {"days": 0, "lines": 0, "movements": 0, "new_ingredients": created, "matched": [
        {"workbook": r["name"], "ingredient": r["ing"]["name"]} for r in resolved.values() if r["name"] not in created]}

    for s in sheets:
        d0, d1 = s["dates"][0], s["dates"][-1]
        day_id = c.execute(
            """INSERT INTO days(business_date, status, opened_at, closed_at, notes, source, covers_to, import_batch_id)
               VALUES (?, 'closed', ?, ?, ?, 'workbook', ?, ?)""",
            (d0, now(), now(), f"Imported from workbook sheet '{s['name']}'" + (f" (covers {d0} to {d1})" if d1 != d0 else ""),
             d1 if d1 != d0 else None, batch)).lastrowid
        stats["days"] += 1
        for row in s["rows"]:
            r = resolved[row["key"]]
            k = r["scale"]
            ing = r["ing"]
            opening, added, closing, system = (round((row[f] or 0) * k, 6) for f in ("opening", "additions", "closing", "system"))
            expected = round(opening + added - system, 6)
            variance = round(closing - expected, 6)
            reason = None
            if row["note"]:
                reason = f"Workbook note: {row['note']}"
            elif abs(variance) > 1e-9:
                reason = "Imported (no reason recorded)"
            c.execute(
                """INSERT INTO day_stock(day_id, ingredient_id, opening, delivered, adjusted, sold, wasted, expected,
                                         counted, variance, reason) VALUES (?,?,?,?,0,?,0,?,?,?,?)""",
                (day_id, ing["id"], opening, added, system, expected, closing, variance, reason))
            stats["lines"] += 1
            note = f"Workbook '{s['name']}'"
            if added:
                c.execute("""INSERT INTO stock_movements(day_id, ingredient_id, type, qty, unit_cost, note, staff_id, created_at)
                             VALUES (?, ?, 'delivery', ?, ?, ?, ?, ?)""", (day_id, ing["id"], added, ing["cost_per_unit"], note, staff_id, now()))
                stats["movements"] += 1
            if system:
                c.execute("""INSERT INTO stock_movements(day_id, ingredient_id, type, qty, unit_cost, note, staff_id, created_at)
                             VALUES (?, ?, 'external_sale', ?, ?, ?, ?, ?)""",
                          (day_id, ing["id"], -system, ing["cost_per_unit"], note + " – other POS", staff_id, now()))
                stats["movements"] += 1
    summary = {k: v for k, v in stats.items()} | {"units": {r["name"]: r["unit"] for r in resolved.values()},
                                                   "warnings": parsed["counts"].get("warning", 0)}
    c.execute("UPDATE import_batches SET summary = ? WHERE id = ?", (json.dumps(summary), batch))
    audit(c, staff_id, "import.stock_history", "import_batch", batch, summary | {"filename": filename})
    return summary
